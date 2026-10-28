"""Build the Claude learning guide (.docx) from chapters/*.md.

Usage: python build_docx.py [output.docx]
Supports the restricted markup used by the chapter files: #/##/### headings,
paragraphs, **bold**, `code`, - bullets, 1. numbered items, ``` code fences,
> TIP:/NOTE:/WARNING:/EXERCISE: callouts and pipe tables.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
CHAPTERS = sorted((HERE / "chapters").glob("*.md"))
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "Learning-Claude-Guide.docx"

ACCENT = RGBColor(0xC1, 0x5F, 0x3C)  # Claude-style terracotta
DARK = RGBColor(0x1F, 0x1F, 0x1F)
CALLOUTS = {
    "TIP": ("Tip", "E8F4EA"),
    "NOTE": ("Note", "E8EEF7"),
    "WARNING": ("Warning", "FBEAE5"),
    "EXERCISE": ("Exercise", "F4EFE6"),
}
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def shade(element, hex_fill):
    pr = element.get_or_add_tcPr() if hasattr(element, "get_or_add_tcPr") else element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def add_field(paragraph, instr):
    run = paragraph.add_run()
    for tag, text in (("begin", None), ("instr", instr), ("separate", None), ("end", None)):
        if tag == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        run._r.append(el)
        if tag == "separate":
            placeholder = OxmlElement("w:t")
            placeholder.text = "Right-click and choose Update Field to build the table of contents." if "TOC" in instr else "1"
            run._r.append(placeholder)


def add_inline(paragraph, text, size=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.color.rgb = RGBColor(0x8B, 0x2E, 0x16)
        else:
            run = paragraph.add_run(part)
        if size:
            run.font.size = size


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, colour in (("Title", 30, ACCENT), ("Heading 1", 22, ACCENT),
                               ("Heading 2", 15, DARK), ("Heading 3", 12.5, ACCENT)):
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = colour
        st.font.bold = name != "Title"
        fonts = st._element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme"):  # drop theme font so Calibri sticks
            fonts.attrib.pop(qn(attr), None)
    styles["Heading 1"].paragraph_format.space_after = Pt(14)
    styles["Heading 2"].paragraph_format.space_before = Pt(16)
    styles["Heading 3"].paragraph_format.space_before = Pt(10)


def setup_page(doc):
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)  # A4
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.3))
    sec.top_margin = sec.bottom_margin = Cm(2.2)


def add_footer(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Learning Claude  |  Page ")
    r.font.size = Pt(9)
    add_field(p, "PAGE")


def title_page(doc):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.add_run("Learning Claude")
    p = doc.add_paragraph()
    r = p.add_run("A Practitioner's Guide to Claude Code and the Claude API")
    r.font.size = Pt(16)
    r.font.color.rgb = DARK
    doc.add_paragraph()
    for line in ("Reference and learning handbook for the Larchmont Hotels capstone team",
                 "Edition: October 2026",
                 "Covers: Claude models, prompting, Claude Code (memory, settings, commands, skills, "
                 "subagents, hooks, MCP, automation) and the Claude API"):
        p = doc.add_paragraph()
        r = p.add_run(line)
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("Product details change quickly. Check the official documentation "
                  "(code.claude.com/docs and platform.claude.com/docs) before relying on a specific "
                  "flag, model ID or price.")
    r.italic = True
    r.font.size = Pt(9.5)
    doc.add_page_break()
    doc.add_heading("Contents", level=1)
    add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')


def code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    shade(cell._tc, "F3F3F3")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for i, line in enumerate(lines):
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(8.5)
        if i < len(lines) - 1:
            run.add_break(WD_BREAK.LINE)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def callout(doc, kind, text):
    label, fill = CALLOUTS[kind]
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    shade(cell._tc, fill)
    p = cell.paragraphs[0]
    r = p.add_run(f"{label}: ")
    r.bold = True
    r.font.color.rgb = ACCENT
    add_inline(p, text)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def data_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c or "--") for c in r)]
    width = max(len(r) for r in cells)
    table = doc.add_table(rows=len(cells), cols=width)
    table.style = "Table Grid"
    for i, row in enumerate(cells):
        for j in range(width):
            cell = table.cell(i, j)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            add_inline(p, row[j] if j < len(row) else "", size=Pt(9))
            if i == 0:
                shade(cell._tc, "F1DED6")
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render(doc, text):
    lines = text.splitlines()
    i, para = 0, []

    def flush():
        if para:
            add_inline(doc.add_paragraph(), " ".join(para))
            para.clear()

    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            flush()
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i].rstrip("\n"))
                i += 1
            code_block(doc, block or [""])
        elif m := re.match(r"^(#{1,3})\s+(.*)", stripped):
            flush()
            level = len(m.group(1))
            if level == 1:
                doc.add_section(WD_SECTION.NEW_PAGE)
            doc.add_heading(m.group(2).replace("**", ""), level=level)
        elif m := re.match(r"^>\s*(TIP|NOTE|WARNING|EXERCISE):\s*(.*)", stripped):
            flush()
            callout(doc, m.group(1), m.group(2))
        elif stripped.startswith(">"):
            flush()
            callout(doc, "NOTE", stripped.lstrip("> ").strip())
        elif stripped.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            data_table(doc, rows)
            continue
        elif m := re.match(r"^[-*]\s+(.*)", stripped):
            flush()
            add_inline(doc.add_paragraph(style="List Bullet"), m.group(1))
        elif m := re.match(r"^(\d+)[.)]\s+(.*)", stripped):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(3)
            p.add_run(f"{m.group(1)}.	")
            add_inline(p, m.group(2))
        elif not stripped:
            flush()
        else:
            para.append(stripped)
        i += 1
    flush()


def main():
    doc = Document()
    setup_styles(doc)
    setup_page(doc)
    add_footer(doc.sections[0])
    # Ask Word to refresh the TOC and page fields when the file is opened.
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    doc.settings.element.append(upd)
    title_page(doc)
    for path in CHAPTERS:
        render(doc, path.read_text(encoding="utf-8"))
    doc.save(OUT)
    print(f"Wrote {OUT} from {len(CHAPTERS)} chapter files")


if __name__ == "__main__":
    main()
